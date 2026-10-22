"""Build NAR_MANUSCRIPT_REVISED_v3.1.docx from NAR_MANUSCRIPT_REVISED.md.

Single source of truth: NAR_MANUSCRIPT_REVISED.md (HEAD). No other .md / .docx is read.
Markdown tables are rendered as real Word tables. Every paragraph of the source
(including all withdrawn / superseded / reproducibility / provenance disclosures)
is carried over verbatim.

Usage:
    python build_docx_v31.py
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

SRC = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(SRC, "NAR_MANUSCRIPT_REVISED.md")
OUT = os.path.join(SRC, "NAR_MANUSCRIPT_REVISED_v3.1.docx")

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
BODY_PT = 12
TABLE_PT = 9

# NAR submission convention: figure legends precede the reference list.
# Set to False to keep the exact section order of the markdown.
MOVE_LEGENDS_BEFORE_REFS = True

HEADING_PT = {1: 16, 2: 14, 3: 13, 4: 12}

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`)")


# --------------------------------------------------------------------------
# low-level helpers
# --------------------------------------------------------------------------
def set_run_font(run, size=BODY_PT, bold=False, italic=False, mono=False):
    run.font.name = MONO_FONT if mono else BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), MONO_FONT if mono else BODY_FONT)


def add_inline(par, text, size=BODY_PT, bold=False, italic=False):
    """Render **bold**, *italic* and `code` spans into runs."""
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            add_inline(par, tok[2:-2], size, True, italic)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            add_inline(par, tok[1:-1], size, bold, True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            set_run_font(par.add_run(tok[1:-1]), size, bold, italic, mono=True)
        else:
            set_run_font(par.add_run(tok), size, bold, italic)


def add_heading(doc, level, text):
    par = doc.add_paragraph()
    par.style = doc.styles["Normal"]
    pf = par.paragraph_format
    pf.space_before = Pt(12 if level <= 2 else 10)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    pf.keep_with_next = True
    if level == 1:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(par, text, size=HEADING_PT[level], bold=True)
    return par


def add_body(doc, text, style_name=None):
    par = doc.add_paragraph()
    if style_name:
        try:
            par.style = doc.styles[style_name]
        except KeyError:
            style_name = None
    par.paragraph_format.space_after = Pt(6)
    par.paragraph_format.line_spacing = 1.5
    add_inline(par, text)
    return par


def add_caption(doc, text):
    """Table / Figure caption: bold lead-in, 11 pt."""
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(10)
    par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.line_spacing = 1.15
    par.paragraph_format.keep_with_next = True
    add_inline(par, text, size=11)
    return par


def set_cell_text(cell, text, bold=False):
    par = cell.paragraphs[0]
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.space_before = Pt(0)
    par.paragraph_format.line_spacing = 1.0
    add_inline(par, text.strip(), size=TABLE_PT, bold=bold)


def mark_header_row(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def add_md_table(doc, rows):
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncol)
    try:
        table.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for ri, cells in enumerate(rows):
        row = table.add_row()
        for ci in range(ncol):
            set_cell_text(row.cells[ci], cells[ci] if ci < len(cells) else "",
                          bold=(ri == 0))
        if ri == 0:
            mark_header_row(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


# --------------------------------------------------------------------------
# markdown source
# --------------------------------------------------------------------------
def split_rows(block):
    """Strip the markdown alignment row and split into cells."""
    out = []
    for i, line in enumerate(block):
        if i == 1 and re.match(r"^\|[\s:|-]+\|$", line.strip()):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        out.append(cells)
    return out


def reorder_sections(lines):
    """Move the top-level FIGURE LEGENDS section in front of REFERENCES."""
    if not MOVE_LEGENDS_BEFORE_REFS:
        return lines
    idx = [i for i, l in enumerate(lines) if re.match(r"^##\s", l)]
    if not idx:
        return lines
    bounds = []
    for k, start in enumerate(idx):
        end = idx[k + 1] if k + 1 < len(idx) else len(lines)
        bounds.append((start, end, lines[start][3:].strip()))
    fig = next((b for b in bounds if b[2].upper().startswith("FIGURE LEGEND")), None)
    ref = next((b for b in bounds if b[2].upper().startswith("REFERENCE")), None)
    if not fig or not ref or fig[0] < ref[0]:
        return lines
    seg = lines[fig[0]:fig[1]]
    ref_seg = lines[ref[0]:fig[0]]
    # keep anything after the last '---' of the legend block (closing note) at the end
    cut = max((k for k, l in enumerate(seg) if l.strip() == "---"), default=len(seg) - 1)
    return lines[:ref[0]] + seg[:cut] + ref_seg + seg[cut:]


# --------------------------------------------------------------------------
# main build
# --------------------------------------------------------------------------
def build():
    with open(MD, encoding="utf-8") as fh:
        lines = fh.read().split("\n")
    lines = reorder_sections(lines)

    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_PT)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.54)
    sec.right_margin = Cm(2.54)

    i = 0
    n = len(lines)
    n_tables = 0
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        # markdown table
        if stripped.startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            add_md_table(doc, split_rows(block))
            n_tables += 1
            continue

        # heading
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            add_heading(doc, len(m.group(1)), m.group(2))
            i += 1
            continue

        # bullet list
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            par = doc.add_paragraph()
            try:
                par.style = doc.styles["List Bullet"]
            except KeyError:
                add_inline(par, "\u2022 " + m.group(1))
            else:
                par.paragraph_format.space_after = Pt(3)
                par.paragraph_format.line_spacing = 1.5
                add_inline(par, m.group(1))
            i += 1
            continue

        # numbered list / reference entry: keep the literal number
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            add_body(doc, stripped)
            i += 1
            continue

        add_body(doc, stripped)
        i += 1

    doc.core_properties.title = (
        "Fusang: Tardigrade Edition - K-mer Frequency Vector Alignment-Free "
        "Phylogenetic Inference Resilient to Indel-Rich Sequence Evolution"
    )
    doc.save(OUT)
    return n_tables


if __name__ == "__main__":
    t = build()
    print("wrote", OUT, "tables:", t)
    sys.exit(0)
