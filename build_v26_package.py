"""Regenerate v2.6 DOCX (black headings), rebuild submission folder+ZIP, All_Tables DOCX."""
import os, re, glob, shutil, zipfile, subprocess, tempfile
from docx import Document
from docx.shared import RGBColor

SRC = r"D:\系统发育树项目\Fusang\Fusang-main"
MD = os.path.join(SRC, "NAR_MANUSCRIPT_REVISED.md")
DOCX = os.path.join(SRC, "NAR_MANUSCRIPT_REVISED_v2.6.docx")
SUBMIT_DIR = os.path.join(SRC, "NAR_Submission_v2.6")
PANDOC = r"C:\Users\admin\miniconda3\envs\human_selection\Library\bin\pandoc.exe"

# 1. pandoc MD -> DOCX
r = subprocess.run([PANDOC, MD, "-o", DOCX, "--standalone"], capture_output=True, text=True)
print("pandoc:", r.returncode, r.stderr[:200])

# 2. black headings
doc = Document(DOCX)
fixed = 0
for para in doc.paragraphs:
    sn = para.style.name if para.style else ""
    if sn.startswith(("Heading", "heading")):
        for run in para.runs:
            run.font.color.rgb = RGBColor(0, 0, 0)
            fixed += 1
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                sn = para.style.name if para.style else ""
                if sn.startswith(("Heading", "heading")):
                    for run in para.runs:
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        fixed += 1
doc.save(DOCX)
print(f"black headings: {fixed} runs fixed")

# 3. All_Tables.docx from manuscript tables
with open(MD, encoding='utf-8') as f:
    content = f.read()
tables = re.findall(r'(\*\*Table (\d+)\. .+?\*\*)\n(.*?)(?=\n\n\*\*Table|\n\n##|\Z)', content, re.DOTALL)
tables_dir = os.path.join(SRC, "tables_pdf")
os.makedirs(tables_dir, exist_ok=True)
all_md = "\n\n\\newpage\n\n".join(f"{c.strip()}\n\n{b.strip()}" for c, n, b in tables)
tmp = os.path.join(tempfile.gettempdir(), "all_tables_v26.md")
open(tmp, 'w', encoding='utf-8').write(all_md)
tables_docx = os.path.join(tables_dir, "All_Tables.docx")
r = subprocess.run([PANDOC, tmp, "-o", tables_docx, "--standalone"], capture_output=True, text=True)
print(f"All_Tables.docx: {len(tables)} tables, rc={r.returncode}")
os.remove(tmp)

# 4. submission folder
os.makedirs(SUBMIT_DIR, exist_ok=True)
for old in glob.glob(os.path.join(SUBMIT_DIR, "*")):
    os.remove(old)
files = [DOCX, os.path.join(SRC, "cover_letter.docx"), os.path.join(SRC, "GraphicalAbstract.pdf"), tables_docx]
files += sorted(glob.glob(os.path.join(SRC, "Figure*.pdf")))
files += sorted(glob.glob(os.path.join(SRC, "Supplementary_Figure_*.pdf")))
files.append(os.path.join(SRC, "repro_package_v1.4.zip"))
copied = []
for fp in files:
    if os.path.exists(fp):
        shutil.copy2(fp, os.path.join(SUBMIT_DIR, os.path.basename(fp)))
        copied.append(os.path.basename(fp))
    else:
        print("MISSING:", fp)
print(f"copied {len(copied)} files")

# 5. ZIP
zip_path = os.path.join(SRC, "NAR_Submission_v2.6.zip")
if os.path.exists(zip_path):
    os.remove(zip_path)
with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
    for fn in sorted(copied):
        zf.write(os.path.join(SUBMIT_DIR, fn), fn)
print(f"ZIP: {zip_path} ({os.path.getsize(zip_path):,} bytes, {len(copied)} files)")
