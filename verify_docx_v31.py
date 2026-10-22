"""Self-check for NAR_MANUSCRIPT_REVISED_v3.1.docx.

Reopens the generated docx, reports structure statistics, and greps for
current (must be present) and superseded (must be judged case by case) values.
"""
import os
import re
import sys

from docx import Document

SRC = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(SRC, "NAR_MANUSCRIPT_REVISED_v3.1.docx")
MD = os.path.join(SRC, "NAR_MANUSCRIPT_REVISED.md")

MUST_PRESENT = ["0.0507", "2.07\u00d710\u207b\u00b9\u2074", "0.669", "0.2617",
                "0.2726", "1.54\u00d710\u207b\u2074", "0.294"]
WATCH = ["p=0.049", "2.3\u00d710\u207b\u00b9\u2074", "0.709", "0.083 \u00b1 0.005",
         "regardless of scale", "indel rates above 0.01"]


def full_text(doc):
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def main():
    doc = Document(DOCX)
    text = full_text(doc)

    n_par = len(doc.paragraphs)
    n_tab = len(doc.tables)
    rows = [len(t.rows) for t in doc.tables]
    n_rows = sum(rows)
    n_sec = len(doc.sections)

    print("=" * 68)
    print("STRUCTURE")
    print("=" * 68)
    print(f"paragraphs (top level)   : {n_par}")
    print(f"tables                   : {n_tab}   (expected 12)")
    print(f"table rows               : {n_rows}")
    print(f"per-table rows           : {rows}")
    print(f"per-table cols           : {[len(t.columns) for t in doc.tables]}")
    print(f"sections                 : {n_sec}")
    print(f"page size                : {doc.sections[0].page_width.cm:.1f} x "
          f"{doc.sections[0].page_height.cm:.1f} cm")

    words = len(re.findall(r"\S+", text))
    size_kb = os.path.getsize(DOCX) / 1024
    print(f"word count               : {words}")
    print(f"file size                : {size_kb:.1f} KB")
    print(f"estimated pages (w/600)  : {words / 600:.1f}")

    print()
    print("=" * 68)
    print("MUST BE PRESENT")
    print("=" * 68)
    for v in MUST_PRESENT:
        c = text.count(v)
        print(f"  [{'OK ' if c else 'MISS'}] {v!r}: {c} occurrence(s)")

    print()
    print("=" * 68)
    print("SUPERSEDED VALUES - CONTEXT INSPECTION")
    print("=" * 68)
    for v in WATCH:
        hits = [m.start() for m in re.finditer(re.escape(v), text)]
        print(f"\n  {v!r}: {len(hits)} hit(s)")
        for h in hits:
            ctx = text[max(0, h - 170):h + 170].replace("\n", " ")
            print(f"     ...{ctx}...")

    print()
    print("=" * 68)
    print("MD <-> DOCX COVERAGE (every non-table source line must appear)")
    print("=" * 68)
    with open(MD, encoding="utf-8") as fh:
        md_lines = fh.read().split("\n")

    def flat(s):
        # stash `code` spans so that '*' inside them is not stripped
        spans = []

        def _stash(m):
            spans.append(m.group(1))
            return "\x00%d\x00" % (len(spans) - 1)

        s = re.sub(r"`([^`]+)`", _stash, s)
        s = re.sub(r"[*`#]", "", s)
        s = re.sub(r"\x00(\d+)\x00", lambda m: spans[int(m.group(1))], s)
        s = re.sub(r"^[\-\u2022]\s+", "", s.strip())
        return re.sub(r"\s+", " ", s).strip()

    flat_docx = re.sub(r"\s+", " ", text)
    md_para = [l.strip() for l in md_lines
               if l.strip() and l.strip() != "---" and not l.strip().startswith("|")]
    missing = [l for l in md_para if flat(l) and flat(l) not in flat_docx]
    print(f"md non-table lines        : {len(md_para)}")
    print(f"non-table lines not found : {len(missing)}")
    for l in missing[:20]:
        print("   MISSING:", l[:160])

    # markdown table cell coverage
    md_cells = []
    for l in md_lines:
        if l.strip().startswith("|") and not re.match(r"^\|[\s:|-]+\|$", l.strip()):
            md_cells += [c.strip() for c in l.strip().strip("|").split("|")]
    cells_missing = [c for c in md_cells if c and flat(c) not in flat_docx]
    print(f"md table cells            : {len(md_cells)}")
    print(f"table cells not found     : {len(cells_missing)}")
    for c in cells_missing[:20]:
        print("   MISSING CELL:", c[:120])

    print()
    print("=" * 68)
    print("REFERENCE NUMBERING (REFERENCES section only)")
    print("=" * 68)
    m = re.search(r"REFERENCES(.*?)(?:FIGURE LEGENDS|Manuscript prepared)", text, re.S)
    seg = m.group(1) if m else ""
    refs = re.findall(r"^\s*(\d+)\.\s+[A-Z]", seg, re.MULTILINE)
    print("reference entries         :", len(refs))
    print("numbering 1..N in order   :",
          refs == [str(k) for k in range(1, len(refs) + 1)])
    print("numbers                   :", refs)

    print()
    print("=" * 68)
    print("SECTION ORDER IN DOCX")
    print("=" * 68)
    order = [p.text.strip() for p in doc.paragraphs
             if re.match(r"^(ABSTRACT|INTRODUCTION|MATERIALS AND METHODS|RESULTS|"
                         r"DISCUSSION|SUPPLEMENTARY MATERIAL|DATA AVAILABILITY|"
                         r"FUNDING|COMPETING INTERESTS|ACKNOWLEDGEMENTS|"
                         r"AUTHOR CONTRIBUTIONS|CORRESPONDING AUTHOR|REFERENCES|"
                         r"FIGURE LEGENDS)$", p.text.strip())]
    print(" -> ".join(order))


if __name__ == "__main__":
    main()
    sys.exit(0)
